"""
Main Execution Script: Generates SA 230 Workpaper & CARO Annexure
"""

from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill
from src.engine import CARO2020Engine


def run_statutory_audit():
    # Test engagement data with realistic audit exceptions
    client_engagement_data = {
        "ppe_title_deed_exceptions": 0,
        "ppe_revaluation_pct": 0.0,
        "loans_prejudicial_to_interest": False,
        "loans_overdue_over_90_days": False,
        "sec_185_186_violations_count": 0,
        "deposits_accepted": False,
        "cost_records_applicable": True,
        "cost_records_maintained": True,
        "unrecorded_income_cr": 0.0,
        "borrowing_defaults_count": 0,
        "wilful_defaulter": False,
        "short_term_for_long_term_cr": 0.0,
        "ipo_or_private_placement_raised": False,
        "fraud_noticed_count": 0,
        "whistleblower_complaints_count": 1,
        "is_nidhi_company": False,
        "rpt_violations_count": 0,
        "internal_audit_adequate": True,
        "sec_192_non_cash_deals": False,
        "fin_asset_pct": 14.0,
        "fin_income_pct": 8.0,
        "cash_loss_cy_cr": 0.0,
        "cash_loss_py_cr": 0.0,
        "auditor_resigned_during_year": False,
        "current_ratio": 1.4,
        "dscr": 2.1,
        "csr_unspent_transferred": True,
        "is_standalone": True,
        # Exceptions flagged:
        "sanctioned_wc_limit_cr": 60.0,
        "quarterly_bank_variances_count": 1,
        "undisputed_dues_over_6m_cr": 0.0,
        "disputed_dues_cr": 24.50,
    }

    client_name = "Apex Consumer Products Limited"
    print("=" * 75)
    print(f"RUNNING STATUTORY AUDIT FOR: {client_name.upper()}")
    print("STANDARDS: SA 230, SA 500 | MANDATE: CARO 2020 (ALL 21 CLAUSES)")
    print("=" * 75)

    engine = CARO2020Engine(client_name, client_engagement_data)
    results = engine.evaluate_all_clauses()

    for r in results:
        status = "[QUALIFIED / EXCEPTION]" if r.has_exception else "[CLEAN]"
        print(f"Clause {r.clause_num:<7} | {r.clause_title:<35} | {status}")

    out_dir = Path("dist")
    out_dir.mkdir(exist_ok=True)
    wp_path = out_dir / "WP_CARO2020_Master_Workpaper.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "WP.CARO.MasterIndex"

    ws["A1"] = f"ENGAGEMENT: {client_name.upper()} | AUDIT PERIOD: FY 2025-26"
    ws["A1"].font = Font(name="Calibri", size=13, bold=True, color="1F497D")
    ws["A2"] = "AUDIT WORKPAPER REF: WP.CARO.2020 | SA 230 AUDIT DOCUMENTATION"
    ws["A2"].font = Font(name="Calibri", size=10, italic=True)

    headers = ["Clause Ref", "Audit Subject Matter", "Applicable?", "Status", "Auditor Findings / Tick Marks"]
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font

    for row_idx, r in enumerate(results, 5):
        ws.cell(row=row_idx, column=1, value=f"Clause {r.clause_num}")
        ws.cell(row=row_idx, column=2, value=r.clause_title)
        ws.cell(row=row_idx, column=3, value="Yes" if r.is_applicable else "No")
        ws.cell(row=row_idx, column=4, value="QUALIFIED" if r.has_exception else "Clean")
        ws.cell(row=row_idx, column=5, value=r.auditor_finding)

    wb.save(wp_path)
    print(f"\n[1] Generated SA 230 Workpaper: {wp_path}")

    report_path = out_dir / "Draft_CARO_Annexure_Report.txt"
    with open(report_path, "w") as f:
        f.write(f"ANNEXURE 'A' TO THE INDEPENDENT AUDITOR'S REPORT OF {client_name.upper()}\n")
        f.write("Statement on matters specified in paragraphs 3 and 4 of CARO 2020\n")
        f.write("=" * 80 + "\n\n")
        for r in results:
            f.write(f"{r.annexure_text}\n\n")

    print(f"[2] Generated Draft CARO Annexure: {report_path}")
    print("=" * 75)


if __name__ == "__main__":
    run_statutory_audit()
